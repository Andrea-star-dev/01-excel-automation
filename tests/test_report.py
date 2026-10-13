from pathlib import Path

import openpyxl

from excel_automation.loader import load_all_orders
from excel_automation.mapping import load_mapping
from excel_automation.report import build_report

CONFIG_PATH = Path(__file__).resolve().parent.parent / "config" / "column_mapping.yaml"
SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_data"


def test_build_report_creates_three_sheets(tmp_path):
    mapping = load_mapping(CONFIG_PATH)
    orders = load_all_orders(SAMPLE_DIR, mapping)

    output_path = build_report(orders, tmp_path / "report.xlsx")
    assert output_path.exists()

    wb = openpyxl.load_workbook(output_path)
    assert wb.sheetnames == ["주문내역", "일일매출요약", "상품별매출"]


def test_summary_sheet_has_total_row(tmp_path):
    mapping = load_mapping(CONFIG_PATH)
    orders = load_all_orders(SAMPLE_DIR, mapping)

    output_path = build_report(orders, tmp_path / "report.xlsx")
    wb = openpyxl.load_workbook(output_path)
    summary_ws = wb["일일매출요약"]

    last_row = [cell.value for cell in summary_ws[summary_ws.max_row]]
    assert last_row[0] == "합계"


def test_date_columns_use_lowercase_format_codes(tmp_path):
    """대문자 YYYY-MM-DD 서식은 Excel은 눈감아주지만 Numbers 등에서는 깨져서 보인다.
    (이 PR에서 실제 Numbers로 열어보다가 발견한 버그에 대한 회귀 테스트.)
    """
    mapping = load_mapping(CONFIG_PATH)
    orders = load_all_orders(SAMPLE_DIR, mapping)

    output_path = build_report(orders, tmp_path / "report.xlsx")
    wb = openpyxl.load_workbook(output_path)

    order_date_cell = wb["주문내역"]["B2"]
    assert order_date_cell.number_format == "yyyy-mm-dd hh:mm:ss"

    summary_date_cell = wb["일일매출요약"]["A2"]
    assert summary_date_cell.number_format == "yyyy-mm-dd"
