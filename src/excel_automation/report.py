"""서식이 적용된 엑셀 리포트 생성 (헤더 스타일, 통화 서식, 합계 행)."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from excel_automation.aggregate import daily_product_summary, daily_summary

HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TOTAL_FONT = Font(bold=True)
CURRENCY_FORMAT = "#,##0"
CURRENCY_COLUMNS = {"unit_price", "total_amount", "revenue_total"}


def _style_header(ws: Worksheet) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autofit_columns(ws: Worksheet, df: pd.DataFrame) -> None:
    for idx, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), df[col].astype(str).map(len).max() if len(df) else 0) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(width, 40)


def _apply_currency_format(ws: Worksheet, df: pd.DataFrame) -> None:
    for idx, col in enumerate(df.columns, start=1):
        if col in CURRENCY_COLUMNS:
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=idx).number_format = CURRENCY_FORMAT


def _write_sheet(writer: pd.ExcelWriter, df: pd.DataFrame, sheet_name: str) -> Worksheet:
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]
    _style_header(ws)
    _apply_currency_format(ws, df)
    _autofit_columns(ws, df)
    return ws


def _append_total_row(ws: Worksheet, df: pd.DataFrame, numeric_columns: list[str]) -> None:
    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=1, value="합계").font = TOTAL_FONT
    for idx, col in enumerate(df.columns, start=1):
        if col in numeric_columns:
            cell = ws.cell(row=total_row, column=idx, value=float(df[col].sum()))
            cell.font = TOTAL_FONT
            if col in CURRENCY_COLUMNS:
                cell.number_format = CURRENCY_FORMAT


def build_report(orders: pd.DataFrame, output_path: str | Path) -> Path:
    """병합된 주문 데이터로 서식 있는 리포트(xlsx) 생성. 시트 3개: 주문내역, 일일매출요약, 상품별매출."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    summary = daily_summary(orders)
    product_summary = daily_product_summary(orders)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _write_sheet(writer, orders, "주문내역")

        summary_ws = _write_sheet(writer, summary, "일일매출요약")
        _append_total_row(summary_ws, summary, ["order_count", "quantity_total", "revenue_total"])

        _write_sheet(writer, product_summary, "상품별매출")

    return output_path
